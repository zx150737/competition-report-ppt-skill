#!/usr/bin/env python3
"""验证 pptx 满足 L2 级原生可编辑。退出码非 0 表示不合格。"""
import sys, argparse
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def verify(path):
    prs = Presentation(path)
    errors = []
    stats = {"slides": len(prs.slides), "textboxes": 0, "autoshapes": 0,
             "text_runs": 0, "pictures": 0, "groups": 0, "charts": 0, "tables": 0}
    for i, slide in enumerate(prs.slides, 1):
        if not slide.shapes:
            errors.append(f"第 {i} 页没有任何 shape")
            continue
        has_text = False
        for shp in slide.shapes:
            st = shp.shape_type
            if st == MSO_SHAPE_TYPE.PICTURE:
                stats["pictures"] += 1
                errors.append(f"第 {i} 页含图片 shape（{shp.name}），破坏可编辑性")
                continue
            if st == MSO_SHAPE_TYPE.GROUP:
                stats["groups"] += 1
                errors.append(f"第 {i} 页含 group shape（{shp.name}），建议拆开")
                continue
            if st == MSO_SHAPE_TYPE.TEXT_BOX: stats["textboxes"] += 1
            elif st == MSO_SHAPE_TYPE.AUTO_SHAPE: stats["autoshapes"] += 1
            elif shp.has_chart: stats["charts"] += 1
            elif shp.has_table: stats["tables"] += 1
            if shp.has_text_frame:
                for para in shp.text_frame.paragraphs:
                    for run in para.runs:
                        if run.text.strip():
                            stats["text_runs"] += 1
                            has_text = True
        if not has_text:
            errors.append(f"第 {i} 页没有任何可编辑文字")
    if stats["pictures"] > 0:
        errors.append(f"共 {stats['pictures']} 个图片 shape —— 违反 L2 可编辑要求")
    if stats["text_runs"] == 0:
        errors.append("整个 deck 无真实文本 run —— 极可能图片拼的")
    if stats["textboxes"] == 0:
        errors.append("无任何文本框 —— 检查是否全用 autoshape 塞文字")
    print("统计：", stats)
    for e in errors: print(f"ERROR {e}")
    print(f"\n{'FAIL' if errors else 'PASS'}  可编辑性检查")
    return 0 if not errors else 1

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    args = ap.parse_args()
    sys.exit(verify(args.pptx))
