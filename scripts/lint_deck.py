#!/usr/bin/env python3
"""校验 outline 是否符合场景纪律。退出码非 0 表示不合格。"""
import json, re, sys, argparse

PAGE_RANGE = {"pitch": (23, 27), "report": (8, 20)}
MAX_BULLETS, MAX_UNITS, MAX_KPI = 6, 28, 4

def units(text):
    return len(re.findall(r"[\u4e00-\u9fff]", text)) + len(re.findall(r"[A-Za-z0-9]+", text))

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("outline"); ap.add_argument("pptx", nargs="?")
    args = ap.parse_args()
    data = json.load(open(args.outline, encoding="utf-8"))
    errors, warnings = [], []
    scenario = data.get("scenario")
    slides = data.get("slides", [])
    if scenario not in PAGE_RANGE:
        errors.append(f"scenario 必须是 pitch 或 report，当前: {scenario}")
        scenario = "pitch"
    lo, hi = PAGE_RANGE[scenario]
    if not (lo <= len(slides) <= hi):
        errors.append(f"页数 {len(slides)} 超出 {scenario} 场景要求 {lo}-{hi}")
    for i, s in enumerate(slides, 1):
        role = s.get("role")
        if "title" not in s or not str(s["title"]).strip():
            errors.append(f"第 {i} 页缺少 title")
        if role == "bullets":
            bs = s.get("bullets", [])
            if len(bs) > MAX_BULLETS:
                errors.append(f"第 {i} 页 bullet {len(bs)} 条，超过 {MAX_BULLETS}")
            for j, b in enumerate(bs, 1):
                if units(b) > MAX_UNITS:
                    errors.append(f"第 {i} 页第 {j} 条 {units(b)} 字元，超过 {MAX_UNITS}: {b}")
                if len(b) > 40:
                    warnings.append(f"第 {i} 页第 {j} 条偏长，建议拆句")
        if role == "kpi" and len(s.get("items", [])) > MAX_KPI:
            errors.append(f"第 {i} 页 KPI {len(s['items'])} 个，超过 {MAX_KPI}")
    if slides:
        if slides[0].get("role") != "cover":
            errors.append("第 1 页必须是 cover")
        if slides[-1].get("role") == "section":
            errors.append("最后一页不能是 section")
    if args.pptx:
        try:
            from pptx import Presentation
            prs = Presentation(args.pptx)
            if len(prs.slides) != len(slides):
                errors.append(f"pptx 页数 {len(prs.slides)} 与 outline {len(slides)} 不一致")
        except ImportError:
            warnings.append("未安装 python-pptx，跳过 pptx 校验")
    for w in warnings: print(f"WARN  {w}")
    for e in errors: print(f"ERROR {e}")
    print(f"\n{'FAIL' if errors else 'PASS'}  {len(slides)} 页  {len(errors)} errors / {len(warnings)} warnings")
    sys.exit(1 if errors else 0)

if __name__ == "__main__":
    main()
